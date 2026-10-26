#!/usr/bin/env python3
"""Generate a repository-grounded mini-project report for RippleRisk."""

from __future__ import annotations

from pathlib import Path
from textwrap import dedent

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Mm
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
ASSET_DIR = ROOT / "report_assets"
OUTPUT_DOCX = ROOT / "Java_Mini_Project_Report.docx"

PROJECT_TITLE = "RippleRisk: Systemic Risk Simulation and Financial Contagion Analysis Platform"


def _font(size: int = 14, bold: bool = False):
    try:
        return ImageFont.truetype("DejaVuSans.ttf", size=size)
    except OSError:
        return ImageFont.load_default()


def make_canvas(name: str, width: int = 1800, height: int = 1000):
    ASSET_DIR.mkdir(exist_ok=True)
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    return img, draw, ASSET_DIR / name


def box(draw: ImageDraw.ImageDraw, xy, text: str, fill="#eef4ff", outline="#1f3b67", text_size=28):
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=3)
    x1, y1, x2, y2 = xy
    font = _font(text_size, bold=True)
    lines = text.split("\n")
    tw = max(draw.textlength(line, font=font) for line in lines)
    draw.multiline_text(
        ((x1 + x2) / 2 - tw / 2, (y1 + y2) / 2 - 16 * len(lines)),
        text,
        fill="#111111",
        font=font,
        align="center",
    )


def arrow(draw: ImageDraw.ImageDraw, start, end, color="#223355", width=4):
    draw.line([start, end], fill=color, width=width)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex > sx else -1
        draw.polygon([(ex, ey), (ex - 18 * direction, ey - 10), (ex - 18 * direction, ey + 10)], fill=color)
    else:
        direction = 1 if ey > sy else -1
        draw.polygon([(ex, ey), (ex - 10, ey - 18 * direction), (ex + 10, ey - 18 * direction)], fill=color)


def label(draw: ImageDraw.ImageDraw, at, text: str, size: int = 24):
    draw.text(at, text, fill="#222222", font=_font(size))


def generate_assets() -> dict[str, Path]:
    assets = {}

    img, draw, path = make_canvas("architecture.png")
    label(draw, (60, 40), "Architecture Diagram: RippleRisk Frontend-Backend Setup", 34)
    box(draw, (120, 220, 620, 420), "React 19 + Vite\nFrontend")
    box(draw, (720, 220, 1200, 420), "Spring Boot\nREST API")
    box(draw, (1320, 220, 1710, 420), "H2 In-Memory\nDatabase")
    box(draw, (720, 600, 1200, 810), "Spring Security\nCORS + Stateless")
    arrow(draw, (620, 320), (720, 320))
    arrow(draw, (1200, 320), (1320, 320))
    arrow(draw, (960, 420), (960, 600))
    label(draw, (640, 280), "HTTP /api")
    label(draw, (1215, 280), "JPA")
    label(draw, (980, 500), "Auth & Access Rules")
    img.save(path)
    assets["architecture"] = path

    img, draw, path = make_canvas("use_case.png")
    label(draw, (60, 40), "UML Use-Case Diagram (Implemented + UI-Only Flows)", 34)
    draw.ellipse((80, 300, 260, 520), outline="#333", width=4)
    label(draw, (120, 390), "User", 30)
    draw.rectangle((360, 180, 1650, 880), outline="#1f3b67", width=4)
    label(draw, (390, 200), "RippleRisk System", 30)
    use_cases = [
        (520, 300, "Login (API-backed)"),
        (520, 460, "Manage Entities (API-backed)"),
        (520, 620, "Manage Exposures (API-backed)"),
        (1130, 300, "Configure Scenario\n(UI state)"),
        (1130, 460, "View Results\n(sample data)"),
        (1130, 620, "Audit/Report\n(sample data)")
    ]
    for x, y, t in use_cases:
        draw.ellipse((x, y, x + 360, y + 110), outline="#2554a1", width=3, fill="#eef4ff")
        draw.multiline_text((x + 25, y + 34), t, fill="#111", font=_font(24))
        arrow(draw, (260, 410), (x, y + 55), width=3)
    img.save(path)
    assets["use_case"] = path

    img, draw, path = make_canvas("class_diagram.png")
    label(draw, (60, 40), "UML Class Diagram (Backend Core Classes)", 34)
    box(draw, (120, 220, 640, 520), "RiskEntity\n- id: Long\n- name: String\n- type: EntityType\n- balance: Double\n- status: EntityStatus", text_size=24)
    box(draw, (820, 220, 1340, 520), "Exposure\n- id: Long\n- fromEntity: RiskEntity\n- toEntity: RiskEntity\n- amount: Double", text_size=24)
    box(draw, (520, 620, 940, 900), "Controllers\nAuthController\nEntityController\nExposureController", text_size=24)
    arrow(draw, (640, 360), (820, 360))
    arrow(draw, (820, 390), (640, 390))
    arrow(draw, (720, 620), (600, 520))
    arrow(draw, (760, 620), (1060, 520))
    label(draw, (690, 330), "Many-to-One links")
    img.save(path)
    assets["class"] = path

    img, draw, path = make_canvas("sequence_login.png")
    label(draw, (60, 40), "Sequence Diagram: Login Workflow", 34)
    participants = [(180, "Browser"), (720, "Frontend Login.tsx"), (1260, "AuthController")]
    for x, name in participants:
        box(draw, (x - 120, 120, x + 120, 200), name, text_size=22)
        draw.line((x, 200, x, 900), fill="#777", width=2)
    arrow(draw, (180, 280), (720, 280))
    label(draw, (300, 245), "Enter email/password")
    arrow(draw, (720, 380), (1260, 380))
    label(draw, (840, 345), "POST /api/auth/login")
    arrow(draw, (1260, 500), (720, 500))
    label(draw, (860, 465), "200 token / 401 error")
    arrow(draw, (720, 620), (180, 620))
    label(draw, (300, 585), "Store token + redirect")
    img.save(path)
    assets["sequence"] = path

    img, draw, path = make_canvas("activity_simulation.png")
    label(draw, (60, 40), "Activity Diagram: Current Simulation User Flow", 34)
    box(draw, (730, 140, 1070, 240), "Start")
    box(draw, (640, 300, 1160, 420), "Configure scenario values\nin Scenarios.tsx")
    box(draw, (640, 490, 1160, 610), "Click Run Simulation\n(UI trigger)")
    box(draw, (640, 680, 1160, 800), "View pre-defined sample\nresults in Results.tsx")
    box(draw, (730, 860, 1070, 960), "End")
    arrow(draw, (900, 240), (900, 300))
    arrow(draw, (900, 420), (900, 490))
    arrow(draw, (900, 610), (900, 680))
    arrow(draw, (900, 800), (900, 860))
    label(draw, (1180, 710), "No backend contagion\nengine in current codebase", 24)
    img.save(path)
    assets["activity"] = path

    img, draw, path = make_canvas("er_diagram.png")
    label(draw, (60, 40), "ER Diagram: H2/JPA Schema", 34)
    box(draw, (240, 280, 760, 680), "entities\nPK id\nname\ntype\nbalance\nstatus", text_size=26)
    box(draw, (1020, 280, 1560, 680), "exposures\nPK id\nfrom_entity_id (FK)\nto_entity_id (FK)\namount", text_size=26)
    arrow(draw, (760, 420), (1020, 420))
    arrow(draw, (760, 540), (1020, 540))
    label(draw, (790, 390), "from_entity_id -> entities.id")
    label(draw, (790, 510), "to_entity_id -> entities.id")
    img.save(path)
    assets["er"] = path

    img, draw, path = make_canvas("college_logo_placeholder.png", width=600, height=300)
    draw.rectangle((20, 20, 580, 280), outline="#777", width=3)
    label(draw, (150, 120), "College Logo Placeholder", 26)
    img.save(path)
    assets["logo_placeholder"] = path

    return assets


def set_page_layout(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1)
    section.different_first_page_header_footer = True


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(8)

    heading1 = document.styles["Heading 1"]
    heading1.font.name = "Times New Roman"
    heading1.font.size = Pt(16)
    heading1.font.bold = True

    heading2 = document.styles["Heading 2"]
    heading2.font.name = "Times New Roman"
    heading2.font.size = Pt(14)
    heading2.font.bold = True


def add_header_footer(document: Document) -> None:
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.text = "RippleRisk Mini Project Report"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Page ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)


def add_para(document: Document, text: str):
    p = document.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def add_chapter(document: Document, title: str):
    document.add_page_break()
    document.add_heading(title, level=1)


def add_toc_field(document: Document):
    document.add_heading("Table of Contents", level=1)
    p = document.add_paragraph()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
    p._p.append(fld)
    add_para(
        document,
        "If the table of contents does not auto-populate, open the file in Microsoft Word, right-click inside the TOC, and choose 'Update Field' then 'Update entire table'.",
    )


def add_code_block(document: Document, code_text: str):
    p = document.add_paragraph(code_text)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Inches(0.3)
    for run in p.runs:
        run.font.name = "Consolas"
        run.font.size = Pt(10)


def add_figure(document: Document, path: Path, caption: str, figure_entries: list[str]):
    document.add_picture(str(path), width=Inches(6.2))
    cap = document.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    figure_entries.append(caption)


def add_table_with_caption(document: Document, caption: str, headers: list[str], rows: list[list[str]], table_entries: list[str]):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for r in rows:
        cells = table.add_row().cells
        for i, value in enumerate(r):
            cells[i].text = value
    table_entries.append(caption)
    cap = document.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True


def build_document(assets: dict[str, Path]) -> None:
    doc = Document()
    set_page_layout(doc)
    configure_styles(doc)
    add_header_footer(doc)

    figures: list[str] = []
    tables: list[str] = []

    # Cover page
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("ENGINEERING MINI PROJECT REPORT\n").bold = True
    p.runs[0].font.size = Pt(18)
    title_run = p.add_run(f"\n{PROJECT_TITLE}\n")
    title_run.bold = True
    title_run.font.size = Pt(16)

    doc.add_paragraph().alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(str(assets["logo_placeholder"]), width=Inches(2.4))
    lp = doc.paragraphs[-1]
    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    cover_lines = [
        "\nSubmitted to: [COLLEGE NAME PLACEHOLDER]",
        "Department: [DEPARTMENT PLACEHOLDER]",
        "Subject: [SUBJECT NAME PLACEHOLDER]",
        "\nSubmitted by:",
        "1. [STUDENT NAME] - [ROLL NUMBER]",
        "2. [STUDENT NAME] - [ROLL NUMBER]",
        "\nProject Guide: [GUIDE NAME PLACEHOLDER]",
        "Academic Year: [ACADEMIC YEAR PLACEHOLDER]",
        "Semester: [SEMESTER PLACEHOLDER]",
    ]
    for line in cover_lines:
        cp = doc.add_paragraph(line)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Certificate
    add_chapter(doc, "Certificate")
    add_para(doc, "This is to certify that the mini project titled 'RippleRisk: Systemic Risk Simulation and Financial Contagion Analysis Platform' is a bonafide work carried out by the students listed on the cover page under the guidance of [GUIDE NAME PLACEHOLDER] in partial fulfillment of the requirements of [SUBJECT NAME PLACEHOLDER].")
    add_para(doc, "College: [COLLEGE NAME PLACEHOLDER]    Department: [DEPARTMENT PLACEHOLDER]    Academic Year: [ACADEMIC YEAR PLACEHOLDER]")
    add_para(doc, "Guide Signature: ____________________    HOD Signature: ____________________    Principal Signature: ____________________")

    # Declaration
    add_chapter(doc, "Student Declaration")
    add_para(doc, "We hereby declare that this report is our original work, prepared using the repository artifacts available in Jinnotgin3/ripple-risk, and has not been submitted previously for any degree or diploma. We acknowledge all referenced sources.")
    add_para(doc, "Student Signatures:\n1. ____________________\n2. ____________________")

    # Acknowledgement
    add_chapter(doc, "Acknowledgement")
    add_para(doc, "We thank [COLLEGE NAME PLACEHOLDER], [DEPARTMENT PLACEHOLDER], and our guide [GUIDE NAME PLACEHOLDER] for continuous support. We also acknowledge the maintainers of the open-source technologies used in this project, including Spring Boot, React, Vite, and related libraries.")

    # Abstract
    add_chapter(doc, "Abstract")
    add_para(doc, "RippleRisk is a full-stack prototype for studying systemic financial risk and contagion behavior in interconnected institutions. The repository combines a Java Spring Boot backend and a React TypeScript frontend. The backend exposes REST APIs for authentication, entity management, and exposure management, with persistence through Spring Data JPA and an in-memory H2 database. Security configuration includes stateless session policy, CORS control, and endpoint authorization rules. The frontend provides routed pages for login, network setup, scenario configuration, result visualization, intervention comparison, executive reporting, and audit viewing. Repository analysis shows that core data-management APIs are implemented and can be integrated with UI flows, while several simulation and reporting views currently use static sample datasets for demonstration. This distinction is important for engineering evaluation: existing code demonstrates architecture, componentization, styling consistency, and dashboard interaction design, but does not yet include a backend contagion computation engine or complete production authentication. The project is therefore suitable as an academic mini project for software architecture, API design, and financial risk dashboard prototyping, with clear future scope for algorithmic simulation, persistent storage migration, robust JWT validation, role-based access control, and automated tests.")

    # TOC + lists
    add_chapter(doc, "Preliminary Pages")
    add_toc_field(doc)
    doc.add_heading("List of Figures", level=1)
    lof_placeholder = doc.add_paragraph("(Populated after figure insertion; see updated list below.)")
    doc.add_heading("List of Tables", level=1)
    lot_placeholder = doc.add_paragraph("(Populated after table insertion; see updated list below.)")

    # Chapter 1
    add_chapter(doc, "Chapter 1: Introduction")
    doc.add_heading("1.1 Background", level=2)
    add_para(doc, "RippleRisk targets systemic-risk modeling where institutions are connected through directed credit/exposure links. The root README and frontend modules describe lender-borrower style networks and shock-oriented analysis workflows.")
    doc.add_heading("1.2 Motivation", level=2)
    add_para(doc, "Traditional isolated risk views can miss cascade effects across counterparty links. This project motivates a connected-network perspective where distress propagation can be observed over multiple rounds.")
    doc.add_heading("1.3 Problem Statement", level=2)
    add_para(doc, "Design and implement a web-based platform that records financial entities and exposures, supports scenario inputs, and presents contagion-oriented analytical views for learning and experimentation.")
    doc.add_heading("1.4 Aim and Objectives", level=2)
    add_para(doc, "Objectives implemented in code include: (i) secure login endpoint and frontend login flow, (ii) CRUD-style entity/exposure APIs, (iii) routed UI modules for scenarios, results, compare, report, and audit pages, and (iv) reusable formatting and UI component patterns.")
    doc.add_heading("1.5 Scope and Significance", level=2)
    add_para(doc, "Current scope is prototype-level. Backend simulation computation and persistent production storage are outside current implementation but are clear extension points.")
    doc.add_heading("1.6 Proposed Solution", level=2)
    add_para(doc, "The repository adopts a decoupled frontend-backend architecture with API proxying during development. This supports modular growth and staged feature completion.")
    add_figure(doc, assets["architecture"], "Figure 1.1: RippleRisk High-Level Architecture", figures)

    # Chapter 2
    add_chapter(doc, "Chapter 2: Literature Survey and Existing System")
    doc.add_heading("2.1 Relevant Technical Sources", level=2)
    add_para(doc, "The implementation aligns with established practices in Spring Boot web development [1], Spring Security [2], Spring Data JPA [3], React [4], and Vite development proxying [5].")
    doc.add_heading("2.2 Existing System in Repository", level=2)
    add_para(doc, "Implemented backend controllers expose endpoints for login, entity management, and exposure management. Frontend pages for scenarios, results, compare, report, and audit presently display static/demo content rather than backend-computed simulation outputs.")
    doc.add_heading("2.3 Limitations", level=2)
    add_para(doc, "A complete contagion algorithm service, persistent production database integration, and robust JWT validation filter chain are not currently present in this codebase.")
    doc.add_heading("2.4 Proposed System Advantages", level=2)
    add_para(doc, "Despite prototype gaps, the repository already demonstrates modular code organization, API boundary clarity, and a clean basis for implementing future quantitative models and stronger security hardening.")

    # Chapter 3
    add_chapter(doc, "Chapter 3: Software Requirements and System Analysis")
    doc.add_heading("3.1 Functional Requirements", level=2)
    add_table_with_caption(
        doc,
        "Table 3.1: Functional Requirements Grounded in Repository",
        ["ID", "Requirement", "Repository Evidence"],
        [
            ["FR1", "User login request and token response", "backend/controller/AuthController.java + frontend/pages/Login.tsx"],
            ["FR2", "Create/read/update/delete risk entities", "backend/controller/EntityController.java"],
            ["FR3", "Create/read/delete exposures", "backend/controller/ExposureController.java"],
            ["FR4", "Scenario configuration UI", "frontend/pages/Scenarios.tsx"],
            ["FR5", "Results/compare/report/audit views", "frontend/pages/Results.tsx, Compare.tsx, Report.tsx, Audit.tsx"],
        ],
        tables,
    )
    doc.add_heading("3.2 Non-Functional Requirements", level=2)
    add_para(doc, "Development-time usability (Vite), modular maintainability (componentized frontend + layered backend), and secure defaults (stateless policy + protected endpoints) are reflected in current code.")
    doc.add_heading("3.3 Hardware and Software Requirements", level=2)
    add_table_with_caption(
        doc,
        "Table 3.2: Execution Requirements",
        ["Category", "Requirement"],
        [["Backend", "Java 21, Maven, Spring Boot 3.3.5"], ["Frontend", "Node.js/npm, React 19, Vite"], ["Database", "H2 in-memory"], ["OS", "Any OS with Java + Node runtime support"]],
        tables,
    )
    doc.add_heading("3.4 Feasibility and Use Cases", level=2)
    add_para(doc, "Technical feasibility is high for academic deployment because dependencies are standard and setup is straightforward. Use cases include authenticated access, entity/exposure maintenance, and scenario/result exploration.")
    add_figure(doc, assets["use_case"], "Figure 3.1: Use-Case Diagram", figures)

    # Chapter 4
    add_chapter(doc, "Chapter 4: System Design and Architecture")
    doc.add_heading("4.1 Module View", level=2)
    add_para(doc, "Frontend modules are organized under pages/components/lib/styles, while backend modules are structured as config/controller/entity/repository.")
    doc.add_heading("4.2 UML/Class Relationships", level=2)
    add_figure(doc, assets["class"], "Figure 4.1: Backend Class Diagram", figures)
    doc.add_heading("4.3 Sequence Workflow", level=2)
    add_figure(doc, assets["sequence"], "Figure 4.2: Login Sequence", figures)
    doc.add_heading("4.4 Activity Flow", level=2)
    add_figure(doc, assets["activity"], "Figure 4.3: Current Simulation Activity Flow", figures)
    doc.add_heading("4.5 ER/Schema Design", level=2)
    add_figure(doc, assets["er"], "Figure 4.4: H2/JPA Schema Diagram", figures)

    # Chapter 5
    add_chapter(doc, "Chapter 5: Technology Stack and Implementation")
    doc.add_heading("5.1 Actual Technology Stack", level=2)
    add_para(doc, "Backend: Spring Boot Web, Data JPA, Validation, Security, H2. Frontend: React, TypeScript, Vite, Tailwind CSS, Lucide icons. Build tools: Maven and npm.")
    doc.add_heading("5.2 Repository Structure", level=2)
    add_code_block(
        doc,
        dedent(
            """
            backend/
              src/main/java/com/ripplerisk/{config,controller,entity,repository}
              src/main/resources/application.properties
            frontend/
              src/{pages,components,lib,styles}
              package.json, vite.config.ts
            """
        ).strip(),
    )
    doc.add_heading("5.3 Important Implemented Workflows", level=2)
    add_para(doc, "Login flow posts credentials from Login.tsx to /api/auth/login and stores token client-side. Entity and exposure management are supported by dedicated REST controllers and JPA repositories.")
    doc.add_heading("5.4 Setup and Run Steps", level=2)
    add_code_block(
        doc,
        "cd backend && mvn clean spring-boot:run\ncd frontend && npm install && npm run dev"
    )
    doc.add_heading("5.5 Sanitized Code Snippets", level=2)
    add_code_block(
        doc,
        "@PostMapping(\"/login\")\npublic ResponseEntity<?> login(@RequestBody Map<String,String> body) { ... }"
    )
    add_code_block(
        doc,
        "const SHOCK_TYPES = [\"CREDIT_DEFAULT\",\"LIQUIDITY_CRUNCH\",\"MARKET_SHOCK\",\"REGULATORY_ACTION\"];"
    )

    # Chapter 6
    add_chapter(doc, "Chapter 6: Testing and Results")
    doc.add_heading("6.1 Available Test Infrastructure", level=2)
    add_para(doc, "No dedicated backend or frontend test classes/spec files were found in the repository snapshot inspected for this report. Therefore, validation focused on build/lint execution and report-generation checks.")
    doc.add_heading("6.2 Test-Case Table (Derived)", level=2)
    add_table_with_caption(
        doc,
        "Table 6.1: Derived Validation Cases",
        ["TC ID", "Scenario", "Expected Behavior", "Status Basis"],
        [
            ["TC1", "POST /api/auth/login with valid demo credentials", "Returns token payload", "Implemented in AuthController"],
            ["TC2", "GET /api/entities", "Returns entity list", "Implemented in EntityController"],
            ["TC3", "POST /api/exposures", "Persists exposure", "Implemented in ExposureController"],
            ["TC4", "Navigate /results page", "Shows sample metrics and timeline", "Implemented in Results.tsx (static data)"],
        ],
        tables,
    )
    doc.add_heading("6.3 Screenshots", level=2)
    add_para(doc, "Screenshot placeholders are used in this report because runtime screenshots were not bundled in the repository. Capture guidance: run backend and frontend locally, open each route, and replace placeholders with actual captures.")
    add_para(doc, "Placeholder list: Login, Networks, Scenarios, Results, Compare, Report, Audit pages.")
    doc.add_heading("6.4 Observed Limitations and Error Handling", level=2)
    add_para(doc, "Because several frontend pages rely on static arrays, simulation outputs are illustrative. Backend validation annotations exist on entity endpoints, and missing IDs return HTTP 404 responses where implemented.")

    # Chapter 7
    add_chapter(doc, "Chapter 7: Conclusion and Future Scope")
    add_para(doc, "RippleRisk demonstrates a clean full-stack baseline for systemic-risk platform development, with concrete API and UI modules already in place. It is suitable for academic demonstration of architecture, REST integration, and financial dashboard design.")
    add_para(doc, "Future scope includes: implementing contagion computation service and persistence of scenario runs, replacing static frontend datasets with API-backed data, introducing strong JWT issuance/validation and role-based authorization, adding automated tests, and migrating from in-memory to persistent database for production readiness.")

    # References
    add_chapter(doc, "References")
    refs = [
        "[1] Spring Boot Reference Documentation, https://docs.spring.io/spring-boot/docs/current/reference/html/",
        "[2] Spring Security Reference, https://docs.spring.io/spring-security/reference/",
        "[3] Spring Data JPA Reference, https://docs.spring.io/spring-data/jpa/reference/",
        "[4] React Documentation, https://react.dev/",
        "[5] Vite Documentation (Server Proxy), https://vite.dev/config/server-options.html#server-proxy",
        "[6] H2 Database Engine Documentation, https://h2database.com/html/main.html",
        "[7] Apache Maven Documentation, https://maven.apache.org/guides/index.html",
    ]
    for r in refs:
        add_para(doc, r)

    # Appendix
    add_chapter(doc, "Appendix")
    doc.add_heading("A. API Samples (Sanitized)", level=2)
    add_code_block(doc, "POST /api/entities\n{\"name\":\"Apex Bank\",\"type\":\"LENDER\",\"balance\":5000000,\"status\":\"HEALTHY\"}")
    add_code_block(doc, "POST /api/exposures\n{\"fromEntity\":{\"id\":1},\"toEntity\":{\"id\":2},\"amount\":800000}")
    doc.add_heading("B. Configuration (Sanitized)", level=2)
    add_code_block(doc, "spring.datasource.url=jdbc:h2:mem:rippledb...\nripple.cors.allowed-origins=http://localhost:5173")
    doc.add_heading("C. Figure Assets Directory", level=2)
    add_para(doc, "All generated diagram assets are stored under report_assets/ and can be edited/re-generated by rerunning generate_report.py.")

    # Update preliminary list placeholders with actual entries
    lof_placeholder.text = "\n".join([f"{i+1}. {v}" for i, v in enumerate(figures)])
    lot_placeholder.text = "\n".join([f"{i+1}. {v}" for i, v in enumerate(tables)])

    doc.save(OUTPUT_DOCX)


def main() -> None:
    assets = generate_assets()
    build_document(assets)
    print(f"Generated: {OUTPUT_DOCX}")
    print(f"Assets: {ASSET_DIR}")


if __name__ == "__main__":
    main()
